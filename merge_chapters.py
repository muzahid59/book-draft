#!/usr/bin/env python3
"""Merge per-chapter .docx files back into the master Bengali AI book."""

import copy
import glob
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

CHAPTERS_DIR = "chapters"
OUTPUT = "এআই_কীভাবে_সিদ্ধান্ত_নেয়_Bengali.docx"
FONT_NAME = "Noto Serif Bengali"

def get_chapter_files():
    pattern = os.path.join(CHAPTERS_DIR, "*.docx")
    files = sorted(glob.glob(pattern))
    if not files:
        raise FileNotFoundError(f"No .docx files found in {CHAPTERS_DIR}/")
    return files

def copy_styles(src_doc, dst_doc):
    src_styles = src_doc.element.find(qn('w:styles'))
    if src_styles is not None:
        dst_styles = dst_doc.element.find(qn('w:styles'))
        if dst_styles is not None:
            dst_styles.getparent().remove(dst_styles)
        dst_doc.element.insert(0, copy.deepcopy(src_styles))

def copy_section_props(src_doc, dst_doc):
    src_section = src_doc.sections[0]
    dst_section = dst_doc.sections[0]
    dst_section.page_width = src_section.page_width
    dst_section.page_height = src_section.page_height
    dst_section.top_margin = src_section.top_margin
    dst_section.bottom_margin = src_section.bottom_margin
    dst_section.left_margin = src_section.left_margin
    dst_section.right_margin = src_section.right_margin

def add_toc(doc):
    body = doc.element.body
    first_para = body.find(qn('w:p'))

    toc_title = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'  <w:pPr>'
        f'    <w:jc w:val="center"/>'
        f'    <w:spacing w:after="240" w:line="360" w:lineRule="auto"/>'
        f'  </w:pPr>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:cs="{FONT_NAME}" w:eastAsia="{FONT_NAME}"/>'
        f'      <w:b/>'
        f'      <w:bCs/>'
        f'      <w:sz w:val="44"/>'
        f'      <w:szCs w:val="44"/>'
        f'    </w:rPr>'
        f'    <w:t>সূচিপত্র</w:t>'
        f'  </w:r>'
        f'</w:p>'
    )
    body.insert(body.index(first_para), toc_title)

    toc_field = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'  <w:pPr>'
        f'    <w:spacing w:after="480"/>'
        f'  </w:pPr>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:cs="{FONT_NAME}"/>'
        f'      <w:sz w:val="24"/><w:szCs w:val="24"/>'
        f'    </w:rPr>'
        f'    <w:fldChar w:fldCharType="begin"/>'
        f'  </w:r>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:cs="{FONT_NAME}"/>'
        f'      <w:sz w:val="24"/><w:szCs w:val="24"/>'
        f'    </w:rPr>'
        f'    <w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText>'
        f'  </w:r>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:cs="{FONT_NAME}"/>'
        f'      <w:sz w:val="24"/><w:szCs w:val="24"/>'
        f'    </w:rPr>'
        f'    <w:fldChar w:fldCharType="separate"/>'
        f'  </w:r>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:cs="{FONT_NAME}"/>'
        f'      <w:sz w:val="24"/><w:szCs w:val="24"/>'
        f'    </w:rPr>'
        f'    <w:t>[Right-click and select "Update Field" to populate the Table of Contents]</w:t>'
        f'  </w:r>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:cs="{FONT_NAME}"/>'
        f'      <w:sz w:val="24"/><w:szCs w:val="24"/>'
        f'    </w:rPr>'
        f'    <w:fldChar w:fldCharType="end"/>'
        f'  </w:r>'
        f'</w:p>'
    )
    body.insert(body.index(first_para), toc_field)

    page_break = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'  <w:r><w:br w:type="page"/></w:r>'
        f'</w:p>'
    )
    body.insert(body.index(first_para), page_break)

def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        for p in footer.paragraphs:
            p.clear()

        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run1 = fp.add_run()
        run1.font.name = FONT_NAME
        run1.font.size = Pt(10)
        run1._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))

        run2 = fp.add_run()
        run2.font.name = FONT_NAME
        run2.font.size = Pt(10)
        run2._element.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>'))

        run3 = fp.add_run()
        run3.font.name = FONT_NAME
        run3.font.size = Pt(10)
        run3._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))

def main():
    chapter_files = get_chapter_files()
    print(f"Found {len(chapter_files)} chapter files:")
    for f in chapter_files:
        print(f"  {f}")

    first_doc = Document(chapter_files[0])
    merged = Document()
    copy_styles(first_doc, merged)
    copy_section_props(first_doc, merged)

    body = merged.element.body
    default_para = body.find(qn('w:p'))
    if default_para is not None:
        body.remove(default_para)

    for file_idx, filepath in enumerate(chapter_files):
        chapter_doc = Document(filepath)

        if file_idx > 0:
            page_break = parse_xml(
                f'<w:p {nsdecls("w")}>'
                f'  <w:r><w:br w:type="page"/></w:r>'
                f'</w:p>'
            )
            body.append(page_break)

        for para in chapter_doc.paragraphs:
            body.append(copy.deepcopy(para._element))

    add_toc(merged)
    add_page_numbers(merged)

    merged.save(OUTPUT)
    print(f"\nMerged into: {OUTPUT}")
    print(f"Total paragraphs: {len(merged.paragraphs)}")

if __name__ == "__main__":
    main()
